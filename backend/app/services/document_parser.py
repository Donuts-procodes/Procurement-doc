from __future__ import annotations

import base64
import io

import re
from PIL import Image
from docx import Document as DocxDocument
from pypdf import PdfReader

from app.schemas.schemas import ImageSpatialAnchor


def extract_text(filename: str, content: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _extract_pdf(content)
    if lower.endswith(".docx"):
        return _extract_docx(content)
    if lower.endswith((".txt", ".md")):
        return content.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {filename}")


def extract_images_with_anchors_pdf(content: bytes, max_pages: int = 10) -> list[ImageSpatialAnchor]:
    """Extract individual distinct embedded images from PDF pages with spatial anchors."""
    anchors: list[ImageSpatialAnchor] = []
    try:
        reader = PdfReader(io.BytesIO(content))
        counter = 0
        for page_num, page in enumerate(reader.pages[:max_pages]):
            if hasattr(page, "images"):
                for img in page.images:
                    img_data = img.data
                    if len(img_data) < 1500:
                        continue
                    try:
                        with Image.open(io.BytesIO(img_data)) as pil_img:
                            w, h = pil_img.size
                            if w < 64 or h < 64:
                                continue
                            aspect_ratio = round(w / max(h, 1), 2)
                    except Exception:
                        w, h, aspect_ratio = 800, 600, 1.33

                    mime = "image/png"
                    if img.name.lower().endswith((".jpg", ".jpeg")):
                        mime = "image/jpeg"
                    elif img.name.lower().endswith(".svg"):
                        mime = "image/svg+xml"
                    b64 = base64.b64encode(img_data).decode("utf-8")
                    counter += 1
                    anchors.append(
                        ImageSpatialAnchor(
                            image_id=f"pdf_asset_{counter}",
                            url_or_base64=f"data:{mime};base64,{b64}",
                            width=w,
                            height=h,
                            aspect_ratio=aspect_ratio,
                            original_page_index=page_num + 1,
                            preceding_heading=f"Page {page_num + 1} Visual Asset",
                            surrounding_text=None,
                        )
                    )
                    if len(anchors) >= 8:
                        break
            if len(anchors) >= 8:
                break
    except Exception:
        pass
    return anchors


def extract_images_from_pdf(content: bytes, max_pages: int = 10) -> list[str]:
    """Extract individual distinct embedded images and diagrams from PDF pages into base64 data URLs."""
    anchors = extract_images_with_anchors_pdf(content, max_pages=max_pages)
    return [a.url_or_base64 for a in anchors]


def extract_images_with_anchors_docx(content: bytes) -> list[ImageSpatialAnchor]:
    """Extract embedded images and diagrams from DOCX with spatial anchors (preceding heading, aspect ratio, dimensions)."""
    anchors: list[ImageSpatialAnchor] = []
    supported_mimes = {"image/png", "image/jpeg", "image/jpg", "image/gif", "image/webp", "image/svg+xml"}
    try:
        doc = DocxDocument(io.BytesIO(content))
        current_heading = "Executive Summary"
        seen_rids: set[str] = set()
        img_counter = 0

        # Pass 1: Walk paragraphs in document order to capture preceding headings and nearby text
        for p_idx, p in enumerate(doc.paragraphs):
            txt = p.text.strip()
            # Detect headings by style or clear title patterns
            if p.style and ("Heading" in p.style.name or "Title" in p.style.name):
                if txt:
                    current_heading = txt
            elif txt and len(txt) < 80 and (txt[0].isdigit() or txt.startswith("Phase") or any(txt.startswith(x) for x in ["Architecture", "System", "Milestone", "Scope", "EVOLUTION", "SLA"])):
                current_heading = txt

            for r in p.runs:
                if "drawing" in r._element.xml:
                    match = re.search(r'r:embed="(rId\d+)"', r._element.xml)
                    if not match:
                        continue
                    rid = match.group(1)
                    if rid in seen_rids:
                        continue
                    rel = doc.part.rels.get(rid)
                    if not rel or "image" not in rel.target_ref:
                        continue

                    img_part = rel.target_part
                    content_type = getattr(img_part, "content_type", "").lower()
                    if content_type not in supported_mimes:
                        continue

                    img_bytes = img_part.blob
                    if len(img_bytes) < 1500:
                        continue

                    try:
                        with Image.open(io.BytesIO(img_bytes)) as pil_img:
                            w, h = pil_img.size
                            if w < 64 or h < 64:
                                continue
                            aspect_ratio = round(w / max(h, 1), 2)
                    except Exception:
                        w, h, aspect_ratio = 800, 600, 1.33

                    # Find nearest preceding non-empty text
                    surrounding = txt
                    if not surrounding and p_idx > 0:
                        for back_idx in range(p_idx - 1, max(-1, p_idx - 5), -1):
                            prev_txt = doc.paragraphs[back_idx].text.strip()
                            if prev_txt:
                                surrounding = prev_txt
                                break

                    seen_rids.add(rid)
                    img_counter += 1
                    b64 = base64.b64encode(img_bytes).decode("utf-8")
                    anchors.append(
                        ImageSpatialAnchor(
                            image_id=f"docx_asset_{img_counter}",
                            url_or_base64=f"data:{content_type};base64,{b64}",
                            width=w,
                            height=h,
                            aspect_ratio=aspect_ratio,
                            preceding_heading=current_heading,
                            surrounding_text=surrounding or None,
                        )
                    )
                    if len(anchors) >= 8:
                        break
            if len(anchors) >= 8:
                break

        # Pass 2: Fallback to any remaining rels not matched in paragraphs (e.g. inside tables or shapes)
        if len(anchors) < 8:
            for rel in doc.part.rels.values():
                if "image" in rel.target_ref:
                    rid = getattr(rel, "rId", "")
                    if rid and rid in seen_rids:
                        continue
                    img_part = rel.target_part
                    content_type = getattr(img_part, "content_type", "").lower()
                    if content_type not in supported_mimes:
                        continue
                    img_bytes = img_part.blob
                    if len(img_bytes) < 1500:
                        continue
                    try:
                        with Image.open(io.BytesIO(img_bytes)) as pil_img:
                            w, h = pil_img.size
                            if w < 64 or h < 64:
                                continue
                            aspect_ratio = round(w / max(h, 1), 2)
                    except Exception:
                        w, h, aspect_ratio = 800, 600, 1.33
                    img_counter += 1
                    b64 = base64.b64encode(img_bytes).decode("utf-8")
                    anchors.append(
                        ImageSpatialAnchor(
                            image_id=f"docx_asset_{img_counter}",
                            url_or_base64=f"data:{content_type};base64,{b64}",
                            width=w,
                            height=h,
                            aspect_ratio=aspect_ratio,
                            preceding_heading="Document Body",
                            surrounding_text=None,
                        )
                    )
                    if len(anchors) >= 8:
                        break
    except Exception:
        pass
    return anchors


def extract_images_from_docx(content: bytes) -> list[str]:
    """Extract embedded images and diagrams from DOCX documents into base64 data URLs."""
    anchors = extract_images_with_anchors_docx(content)
    return [a.url_or_base64 for a in anchors]



def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    doc = DocxDocument(io.BytesIO(content))
    parts: list[str] = []

    # Extract paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    # Extract tables formatted as markdown pipe tables
    for table in doc.tables:
        table_lines: list[str] = []
        for i, row in enumerate(table.rows):
            cell_texts = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            table_lines.append("| " + " | ".join(cell_texts) + " |")
            if i == 0:
                table_lines.append("| " + " | ".join(["---"] * len(cell_texts)) + " |")
        if table_lines:
            parts.append("\n".join(table_lines))

    return "\n\n".join(parts)


def extract_structured_tables_docx(content: bytes) -> list[dict[str, Any]]:
    """Extract structured tabular data (headers, rows) from DOCX tables."""
    tables: list[dict[str, Any]] = []
    try:
        doc = DocxDocument(io.BytesIO(content))
        for table in doc.tables:
            if not table.rows or len(table.rows) < 2:
                continue
            headers = [c.text.replace("\n", " ").strip() for c in table.rows[0].cells]
            rows = []
            for row in table.rows[1:]:
                row_cells = [c.text.replace("\n", " ").strip() for c in row.cells]
                if any(row_cells):
                    rows.append(row_cells)
            if headers and rows:
                tables.append({"headers": headers, "rows": rows})
    except Exception:
        pass
    return tables


def extract_structured_tables_from_text(text: str) -> list[dict[str, Any]]:
    """Extract structured tabular data (headers, rows) from markdown pipe tables in text."""
    tables: list[dict[str, Any]] = []
    if not text:
        return tables

    lines = [l.strip() for l in text.split("\n")]
    current_table_lines: list[str] = []

    def process_block(block: list[str]):
        if len(block) < 3:
            return
        has_divider = any(re.match(r"^\|[\s\-:|]+\|$", l) for l in block)
        if not has_divider:
            return
        clean_lines = [l for l in block if not re.match(r"^\|[\s\-:|]+\|$", l)]
        if len(clean_lines) < 2:
            return

        def split_row(r: str) -> list[str]:
            return [c.strip() for c in r.strip("|").split("|")]

        headers = split_row(clean_lines[0])
        rows = [split_row(r) for r in clean_lines[1:]]
        if headers and rows and any(len(r) == len(headers) for r in rows):
            tables.append({"headers": headers, "rows": [r for r in rows if any(r)]})

    for line in lines:
        if line.startswith("|") and line.endswith("|") and line.count("|") >= 3:
            current_table_lines.append(line)
        else:
            if current_table_lines:
                process_block(current_table_lines)
                current_table_lines = []

    if current_table_lines:
        process_block(current_table_lines)

    return tables

_HEADING_PATTERN = re.compile(
    r'^(?:'
    r'\d+[\.\)]\s+'              # "1. Title" or "1) Title"
    r'|[A-Z][A-Z\s\&]{6,}$'     # "EXECUTIVE SUMMARY" (all-caps, 7+ chars)
    r'|#{1,3}\s+'                # Markdown headings "# Title", "## Title"
    r')',
    re.MULTILINE,
)


def _is_heading(line: str) -> bool:
    """Detect whether a paragraph looks like a section heading."""
    stripped = line.strip()
    if not stripped or len(stripped) > 120:
        return False
    return bool(_HEADING_PATTERN.match(stripped))


def chunk_text(text: str, chunk_size: int = 600, overlap: int = 100) -> list[str]:
    """Section-aware text chunking that preserves semantic boundaries.

    Keeps section headings attached to their body paragraphs so vector
    retrieval returns heading+context together instead of orphaned fragments.
    """
    if not text.strip():
        return []

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    current_chunk: list[str] = []
    current_word_count = 0

    for para in paragraphs:
        para_words = len(para.split())

        # Section-boundary detection: if this paragraph is a heading and we already
        # have accumulated content, flush the current chunk before starting the new section.
        if _is_heading(para) and current_chunk and current_word_count > 50:
            chunk_str = "\n\n".join(current_chunk)
            chunks.append(chunk_str)

            # Overlap: carry over the last few paragraphs for context continuity
            overlap_words = 0
            overlap_chunk: list[str] = []
            for p in reversed(current_chunk):
                p_count = len(p.split())
                if overlap_words + p_count <= overlap:
                    overlap_chunk.insert(0, p)
                    overlap_words += p_count
                else:
                    break
            current_chunk = overlap_chunk
            current_word_count = overlap_words

        # Standard word-count flush (unchanged logic, but now section-aware)
        if current_word_count + para_words > chunk_size and current_chunk:
            # Don't flush if the current paragraph is a heading — keep it with its body
            if not _is_heading(para):
                chunk_str = "\n\n".join(current_chunk)
                chunks.append(chunk_str)

                overlap_words = 0
                overlap_chunk: list[str] = []
                for p in reversed(current_chunk):
                    p_count = len(p.split())
                    if overlap_words + p_count <= overlap:
                        overlap_chunk.insert(0, p)
                        overlap_words += p_count
                    else:
                        break
                current_chunk = overlap_chunk
                current_word_count = overlap_words

        current_chunk.append(para)
        current_word_count += para_words

    if current_chunk:
        chunks.append("\n\n".join(current_chunk))

    return chunks
